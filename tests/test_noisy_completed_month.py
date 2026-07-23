"""Noisy completed month: a filled past month must stay untouched; only the target month is filled.

Uses a fully synthetic fixture (no commercial corpus). Ported to the clean
FillPlan API (explicit plan, no suggest_fill_plan oracle).
"""

from __future__ import annotations

import shutil
import tempfile
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from support.noisy_completed_month_fixture import (  # type: ignore[import-not-found]
    JULY_SOURCE_BASE,
    JUNE_SENTINEL_BASE,
    build_noisy_workspace,
    june_h_fingerprint,
)

from corpclaw_lite.agent.fill_plan import (
    FillPeriodPlan,
    FillPlan,
    FillSheetPlan,
    FillSourceRef,
    apply_fill_plan,
)
from corpclaw_lite.agent.workbook_brief import (
    build_workbook_brief,
    format_files_brief_for_agent,
)
from corpclaw_lite.extensions.tools.builtin.excel_workbook import (
    _eval_simple_date_formula,
    _parse_date_key,
)


@pytest.fixture()
def noisy_ws() -> Path:
    td = Path(tempfile.mkdtemp(prefix="noisy_month_"))
    build_noisy_workspace(td)
    yield td
    shutil.rmtree(td, ignore_errors=True)


def _july_only_plan() -> FillPlan:
    """Explicit region plan: fill Июль from two day sources; never touch Июнь."""
    return FillPlan(
        template="template.xlsx",
        output_path="target_report.xlsx",
        cutoff="2026-07-12",
        sheets=[
            FillSheetPlan(
                sheet="Июль",
                date_column="F",
                value_column="H",
                aggregate="sum",
                match_mode="date",
                period=FillPeriodPlan(cell="F6", start_mode="month_start", format="dmy"),
                sources=[
                    FillSourceRef(
                        file="source_a.xlsx",
                        sheet="Report",
                        date_column="A",
                        value_column="B",
                    ),
                    FillSourceRef(
                        file="source_b.xlsx",
                        sheet="Report",
                        date_column="A",
                        value_column="B",
                    ),
                ],
            )
        ],
    )


def test_noisy_brief_lists_june_and_july(noisy_ws: Path) -> None:
    files = [
        noisy_ws / "template.xlsx",
        noisy_ws / "source_a.xlsx",
        noisy_ws / "source_b.xlsx",
        noisy_ws / "distractor_june.xlsx",
    ]
    brief = build_workbook_brief(files)
    text = format_files_brief_for_agent(brief)
    assert "### TEMPLATE HINTS" in text
    assert "Июнь" in text
    assert "Июль" in text
    assert "### SOURCE REGIONS" in text
    assert "distractor_june" in text or "июнь" in text.casefold()


def test_noisy_explicit_plan_targets_july_only() -> None:
    plan = _july_only_plan()
    sheets = [s.sheet for s in plan.sheets]
    assert sheets == ["Июль"]
    assert all(s.match_mode == "date" for s in plan.sheets)


def test_noisy_apply_fills_july_leaves_june(noisy_ws: Path) -> None:
    before = june_h_fingerprint(noisy_ws / "template.xlsx", "Июнь")
    assert before["H16"] == JUNE_SENTINEL_BASE + 1
    assert before["H45"] == JUNE_SENTINEL_BASE + 30

    results = apply_fill_plan(_july_only_plan(), workspace=noisy_ws)
    assert results
    assert not any("Error" in r for r in results)

    out = noisy_ws / "target_report.xlsx"
    after = june_h_fingerprint(out, "Июнь")
    assert after == before, "completed Июнь sheet was modified"

    wb = openpyxl.load_workbook(out)
    wb_v = openpyxl.load_workbook(out, data_only=True)
    try:
        ws = wb["Июль"]
        ws_v = wb_v["Июль"]
        memo: dict[tuple[int, int], date | None] = {}
        found = None
        for row_idx in range(1, (ws.max_row or 0) + 1):
            raw = ws_v.cell(row=row_idx, column=6).value
            cell_date = _parse_date_key(raw)
            if cell_date is None:
                fr = ws.cell(row=row_idx, column=6).value
                cell_date = _parse_date_key(fr)
                if cell_date is None and isinstance(fr, str) and fr.startswith("="):
                    cell_date = _eval_simple_date_formula(ws, row_idx, 6, memo)
            if cell_date == date(2026, 7, 6):
                found = ws.cell(row=row_idx, column=8).value
                break
        expected_july_6 = JULY_SOURCE_BASE["a"][6] + JULY_SOURCE_BASE["b"][6]
        assert int(found or 0) == expected_july_6
        assert "12.07.2026" in str(ws["F6"].value or "")
        assert "30.06.2026" in str(wb["Июнь"]["F6"].value or "")
        assert "12.07.2026" not in str(wb["Июнь"]["F6"].value or "")
    finally:
        wb.close()
        wb_v.close()
