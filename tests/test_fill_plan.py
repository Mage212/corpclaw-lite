from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from corpclaw_lite.agent.fill_plan import (
    FillPeriodPlan,
    FillPlan,
    FillSheetPlan,
    FillSourceRef,
    apply_fill_plan,
    parse_fill_plan,
)
from corpclaw_lite.agent.workbook_brief import (
    build_workbook_brief,
    extract_date_values,
    format_files_brief_for_agent,
)
from corpclaw_lite.extensions.tools.builtin.excel_workbook import fill_by_date


def _save_date_source(path: Path, values: list[tuple[date, int]]) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws.append(["Date", "Value"])
    for day, value in values:
        ws.append([day, value])
    wb.save(path)


def _save_key_source(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Labels"
    ws.append(["Name", "Value"])
    ws.append(["Alpha", 11])
    ws.append(["Beta", 13])
    wb.save(path)


def _save_template(path: Path) -> None:
    wb = Workbook()
    daily = wb.active
    assert daily is not None
    daily.title = "Daily"
    daily["C1"] = "Reporting period"
    daily["D1"] = "01/03/2030 - 31/03/2030"
    daily.append(["Date", "Value"])
    daily.append([date(2030, 3, 8), None])
    daily.append([date(2030, 3, 9), 999])
    keyed = wb.create_sheet("ByLabel")
    keyed.append(["Name", "Value"])
    keyed.append(["Alpha", None])
    keyed.append(["Beta", None])
    wb.save(path)


def _date_sheet(*sources: str, period: FillPeriodPlan | None = None) -> FillSheetPlan:
    return FillSheetPlan(
        sheet="Daily",
        date_column="A",
        value_column="B",
        sources=[
            FillSourceRef(file=source, sheet="Data", date_column="A", value_column="B")
            for source in sources
        ],
        aggregate="sum" if len(sources) > 1 else "first",
        period=period,
    )


def test_multi_source_multi_sheet_fill_is_explicit_and_only_empty(tmp_path: Path) -> None:
    _save_template(tmp_path / "template.xlsx")
    _save_date_source(tmp_path / "source_a.xlsx", [(date(2030, 3, 8), 100)])
    _save_date_source(tmp_path / "source_b.xlsx", [(date(2030, 3, 8), 7)])
    _save_key_source(tmp_path / "keys.xlsx")
    plan = FillPlan(
        template="template.xlsx",
        output_path="completed.xlsx",
        cutoff="2030-03-09",
        sheets=[
            _date_sheet(
                "source_a.xlsx",
                "source_b.xlsx",
                period=FillPeriodPlan(cell="D1", start_mode="source_min", format="preserve"),
            ),
            FillSheetPlan(
                sheet="ByLabel",
                date_column="A",
                value_column="B",
                sources=[
                    FillSourceRef(
                        file="keys.xlsx",
                        sheet="Labels",
                        date_column="A",
                        value_column="B",
                    )
                ],
                match_mode="key",
            ),
        ],
    )

    results = apply_fill_plan(plan, workspace=tmp_path)

    assert all("Error" not in result for result in results)
    wb = openpyxl.load_workbook(tmp_path / "completed.xlsx", data_only=False)
    assert wb["Daily"]["B3"].value == 107
    assert wb["Daily"]["B4"].value == 999
    assert wb["Daily"]["D1"].value == "08/03/2030 - 09/03/2030"
    assert wb["ByLabel"]["B2"].value == 11
    assert wb["ByLabel"]["B3"].value == 13
    wb.close()


@pytest.mark.parametrize(
    ("start_mode", "period_format", "expected"),
    [
        ("preserve", "preserve", "01/03/2030 - 09/03/2030"),
        ("source_min", "iso", "2030-03-08 - 2030-03-09"),
        ("month_start", "dmy", "01.03.2030 - 09.03.2030"),
    ],
)
def test_period_strategies(
    tmp_path: Path, start_mode: str, period_format: str, expected: str
) -> None:
    _save_template(tmp_path / "template.xlsx")
    _save_date_source(tmp_path / "source.xlsx", [(date(2030, 3, 8), 5)])
    plan = FillPlan(
        template="template.xlsx",
        output_path="completed.xlsx",
        cutoff="2030-03-09",
        sheets=[
            _date_sheet(
                "source.xlsx",
                period=FillPeriodPlan(cell="D1", start_mode=start_mode, format=period_format),
            )
        ],
    )

    apply_fill_plan(plan, workspace=tmp_path)

    wb = openpyxl.load_workbook(tmp_path / "completed.xlsx")
    assert wb["Daily"]["D1"].value == expected
    wb.close()


def test_omitted_period_leaves_cell_unchanged(tmp_path: Path) -> None:
    _save_template(tmp_path / "template.xlsx")
    _save_date_source(tmp_path / "source.xlsx", [(date(2030, 3, 8), 5)])
    plan = FillPlan(
        template="template.xlsx",
        output_path="completed.xlsx",
        cutoff="2030-03-09",
        sheets=[_date_sheet("source.xlsx")],
    )

    apply_fill_plan(plan, workspace=tmp_path)

    wb = openpyxl.load_workbook(tmp_path / "completed.xlsx")
    assert wb["Daily"]["D1"].value == "01/03/2030 - 31/03/2030"
    wb.close()


def test_preserve_period_fails_when_existing_cell_has_no_date(tmp_path: Path) -> None:
    _save_template(tmp_path / "template.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "template.xlsx")
    wb["Daily"]["D1"] = "not a date range"
    wb.save(tmp_path / "template.xlsx")
    wb.close()
    _save_date_source(tmp_path / "source.xlsx", [(date(2030, 3, 8), 5)])
    plan = FillPlan(
        template="template.xlsx",
        output_path="completed.xlsx",
        cutoff="2030-03-09",
        sheets=[
            _date_sheet(
                "source.xlsx",
                period=FillPeriodPlan(cell="D1", start_mode="preserve", format="preserve"),
            )
        ],
    )

    with pytest.raises(ValueError, match="has no date"):
        apply_fill_plan(plan, workspace=tmp_path)


def test_parser_rejects_implicit_or_model_supplied_period_values() -> None:
    base = {
        "template": "template.xlsx",
        "cutoff": "2030-03-09",
        "sheets": [
            {
                "sheet": "Daily",
                "date_column": "A",
                "value_column": "B",
                "sources": [
                    {
                        "file": "source.xlsx",
                        "sheet": "Data",
                        "date_column": "A",
                        "value_column": "B",
                    }
                ],
            }
        ],
    }
    base["sheets"][0]["period_cell"] = "D1"
    with pytest.raises(ValueError, match="obsolete"):
        parse_fill_plan(base)
    del base["sheets"][0]["period_cell"]
    base["sheets"][0]["period_value"] = "invented"
    with pytest.raises(ValueError, match="not allowed"):
        parse_fill_plan(base)


def test_period_requires_cutoff_and_source_min_rejects_key_mode() -> None:
    sheet = {
        "sheet": "ByLabel",
        "date_column": "A",
        "value_column": "B",
        "match_mode": "key",
        "period": {"cell": "D1", "start_mode": "source_min", "format": "iso"},
        "sources": [
            {
                "file": "source.xlsx",
                "sheet": "Data",
                "date_column": "A",
                "value_column": "B",
            }
        ],
    }
    with pytest.raises(ValueError, match="cutoff"):
        parse_fill_plan({"template": "template.xlsx", "sheets": [sheet]})
    with pytest.raises(ValueError, match="not valid for key"):
        parse_fill_plan({"template": "template.xlsx", "cutoff": "2030-03-09", "sheets": [sheet]})


def test_files_brief_contains_structure_but_no_values_or_fill_command(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    target = tmp_path / "target.xlsx"
    _save_date_source(source, [(date(2030, 3, 8), 918273)])
    _save_template(target)

    text = format_files_brief_for_agent(build_workbook_brief([source, target]))

    assert "FILES_BRIEF" in text
    assert "date_range=2030-03-08..2030-03-08" in text
    assert "918273" not in text
    assert "Call apply_fill_plan" not in text
    assert "Choose tools according to the user's request" in text


def test_delayed_header_and_total_row_are_extracted_without_fixed_coordinates(
    tmp_path: Path,
) -> None:
    path = tmp_path / "delayed.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws.append(["Metadata"])
    ws.append([])
    ws.append(["Date", "Value"])
    ws.append([date(2030, 3, 8), 5])
    ws.append(["Total", 5])
    wb.save(path)

    assert extract_date_values(path, sheet="Data", date_column="A", value_column="B") == {
        "2030-03-08": 5
    }


def test_formula_dates_are_filled_without_cached_values(tmp_path: Path) -> None:
    path = tmp_path / "formula_target.xlsx"
    output = tmp_path / "formula_completed.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Daily"
    ws.append(["Date", "Value"])
    ws.append([date(2030, 3, 8), None])
    ws.append(["=A2+1", None])
    wb.save(path)

    result = fill_by_date(
        path,
        "Daily",
        "A",
        "B",
        {"2030-03-08": 5, "2030-03-09": 7},
        output,
    )

    assert result.startswith("SUCCESS:")
    completed = openpyxl.load_workbook(output, data_only=False)
    assert completed["Daily"]["B2"].value == 5
    assert completed["Daily"]["B3"].value == 7
    completed.close()


def test_source_path_cannot_escape_workspace(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    _save_template(workspace / "template.xlsx")
    _save_date_source(tmp_path / "outside.xlsx", [(date(2030, 3, 8), 5)])
    plan = FillPlan(
        template="template.xlsx",
        output_path="completed.xlsx",
        cutoff="2030-03-09",
        sheets=[_date_sheet("../outside.xlsx")],
    )

    with pytest.raises(ValueError, match="Source path rejected"):
        apply_fill_plan(plan, workspace=workspace)
