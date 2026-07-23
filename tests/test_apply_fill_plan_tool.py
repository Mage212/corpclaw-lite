"""Unit tests for ApplyFillPlanTool (region-mapping FillPlan).

Tests the production tool wrapper (ApplyFillPlanTool.execute) and the
low-level helpers it depends on (_parse_cutoff, _is_fill_result_error,
resolve_fill_plan workspace-boundary enforcement).
"""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest

from corpclaw_lite.agent.fill_plan import (
    FillPlan,
    FillSheetPlan,
    _is_fill_result_error,
    _parse_cutoff,
    apply_fill_plan,
)
from corpclaw_lite.extensions.tools.builtin.apply_fill_plan import ApplyFillPlanTool


@pytest.mark.asyncio
async def test_apply_fill_plan_tool_rejects_values_only_plan() -> None:
    """A plan sheet with ``values`` but no ``sources`` must be rejected."""
    tool = ApplyFillPlanTool()
    result = await tool.execute(
        plan=json.dumps(
            {
                "template": "template.xlsx",
                "sheets": [
                    {
                        "sheet": "Июль",
                        "date_column": "F",
                        "value_column": "H",
                        "values": {"2026-07-06": 1},
                    }
                ],
            }
        )
    )
    assert result.startswith("Error:")
    assert "values" in result.casefold() or "sources" in result.casefold()


def test_apply_fill_plan_multi_sheet_continues_on_empty(tmp_path: Path) -> None:
    """CR-6: when a sheet has empty ``values``, apply_fill_plan must ``continue``
    so the remaining sheets still run and report. With ``break`` (the bug) only
    the failed sheet would appear in ``results``.
    """
    import openpyxl

    template = tmp_path / "template.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Июль"
    ws["F12"] = "Дата"
    ws["H12"] = "Факт"
    ws.cell(row=16, column=6, value=date(2026, 7, 1))
    ws.cell(row=17, column=6, value=date(2026, 7, 2))
    wb.save(str(template))
    wb.close()

    plan = FillPlan(
        template=str(template),
        output_path=str(tmp_path / "out.xlsx"),
        cutoff=None,
        sheets=[
            FillSheetPlan(
                sheet="Июль",
                date_column="F",
                value_column="H",
                sources=[],
                values={"2026-07-01": 100, "2026-07-02": 200},
            ),
            FillSheetPlan(
                sheet="Missing",
                date_column="F",
                value_column="H",
                sources=[],
                values={},
            ),
            FillSheetPlan(
                sheet="Июль",
                date_column="F",
                value_column="H",
                sources=[],
                values={"2026-07-01": 300},
            ),
        ],
    )

    results = apply_fill_plan(plan, workspace=tmp_path)

    # All three sheets reported.
    assert len(results) == 3, f"expected 3 results, got {len(results)}: {results}"
    assert results[0].startswith("[Июль]"), results[0]
    assert "Error" in results[1], results[1]
    assert results[2].startswith("[Июль]"), results[2]
    # Critical assertion: with the old ``break``, results[2] would never exist.
    assert "SUCCESS" in results[2] or "NOOP" in results[2] or "wrote" in results[2], results[2]


def test_parse_cutoff_rejects_two_digit_year() -> None:
    """CR-20: ``%d.%m.%y`` (2-digit year) is rejected — reporting footgun.

    ``"12.07.26"`` must NOT silently parse to 2026-07-12; require explicit
    4-digit years via ISO (``%Y-%m-%d``) or ``%d.%m.%Y``.
    """
    assert _parse_cutoff("2026-07-12") == date(2026, 7, 12)
    assert _parse_cutoff("12.07.2026") == date(2026, 7, 12)
    assert _parse_cutoff(None) is None
    # 2-digit year must raise.
    with pytest.raises(ValueError):
        _parse_cutoff("12.07.26")


def test_is_fill_result_error_classifies_correctly() -> None:
    """CR-19: centralized error classifier covers all filler result prefixes."""
    assert _is_fill_result_error("Error: invalid period_cell 'X1': bad")
    assert _is_fill_result_error("Error saving file: disk full")
    # Non-error prefixes must NOT classify as errors.
    assert not _is_fill_result_error("SUCCESS: wrote 3 gap value(s) ...")
    assert not _is_fill_result_error("PARTIAL: wrote 1 gap value(s) ...")
    assert not _is_fill_result_error("NOOP: 0 new values written ...")
    assert not _is_fill_result_error("Dry-run fill_by_date ...")


def test_resolve_fill_plan_rejects_workspace_escape(tmp_path: Path) -> None:
    """CR-7: when workspace is explicit, source paths that escape it are rejected.

    ``resolve_fill_plan`` must route through resolve_and_validate_path so the
    workspace boundary holds even for callers that bypass ApplyFillPlanTool.
    """
    from corpclaw_lite.agent.fill_plan import FillSourceRef, resolve_fill_plan

    outside = tmp_path.parent / "secret_target.xlsx"
    outside.write_bytes(b"PK")

    plan = FillPlan(
        template="template.xlsx",
        output_path="out.xlsx",
        cutoff="2026-07-12",
        sheets=[
            FillSheetPlan(
                sheet="Июль",
                date_column="F",
                value_column="H",
                sources=[
                    FillSourceRef(
                        file=f"../{outside.name}",
                        sheet="Report",
                        date_column="A",
                        value_column="B",
                    )
                ],
            )
        ],
    )

    with pytest.raises(ValueError) as excinfo:
        resolve_fill_plan(plan, workspace=tmp_path)
    assert "Source path rejected" in str(excinfo.value) or "outside of workspace" in str(
        excinfo.value
    ), str(excinfo.value)
