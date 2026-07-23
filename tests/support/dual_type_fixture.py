"""Synthetic dual-type media workspace: daily sheet + creatives sheet, 4 sources.

Used to verify the model fills both sheets in one ``apply_fill_plan`` call and
maps day sources → По_дням, creative sources → Креативы (sum across cabinets).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

__all__ = [
    "CREATIVE_ORACLE",
    "DAY_ORACLE",
    "build_dual_type_workspace",
]

# Expected Fact after sum(cab_a, cab_b), cutoff 2026-07-12.
DAY_ORACLE: dict[str, int] = {
    "2026-07-06": 1100,
    "2026-07-07": 2000,
    "2026-07-08": 3000,
    "2026-07-09": 4000,
    "2026-07-10": 5000,
    "2026-07-11": 6000,
    "2026-07-12": 7000,
}

CREATIVE_ORACLE: dict[str, int] = {
    "Creative_A": 10111,
    "Creative_B": 20222,
    "Creative_C": 5333,
}


def build_dual_type_workspace(workspace: Path) -> dict[str, Any]:
    """Create template + 4 raw reports under ``workspace``."""
    import openpyxl

    workspace.mkdir(parents=True, exist_ok=True)

    _write_day_source(
        workspace / "cab_a_days.xlsx",
        {
            "2026-07-06": 1000,
            "2026-07-07": 2000,
            "2026-07-08": 3000,
            "2026-07-09": 4000,
            # past cutoff — must be ignored when cutoff=2026-07-12
            "2026-07-13": 99999,
        },
        headers=("Дата", "Показы"),
    )
    _write_day_source(
        workspace / "cab_b_days.xlsx",
        {
            "2026-07-06": 100,  # overlaps cab_a → sum
            "2026-07-10": 5000,
            "2026-07-11": 6000,
            "2026-07-12": 7000,
        },
        headers=("Day", "Impressions"),
    )
    _write_creative_source(
        workspace / "cab_a_creatives.xlsx",
        {
            "Creative_A": 10000,
            "Creative_B": 20000,
            "Creative_C": 5000,
            "Creative_Noise": 88888,  # not in template — ignore
        },
        headers=("Креатив", "Показы"),
    )
    _write_creative_source(
        workspace / "cab_b_creatives.xlsx",
        {
            "Creative_A": 111,
            "Creative_B": 222,
            "Creative_C": 333,
        },
        headers=("Creative", "Impressions"),
    )

    tmpl = workspace / "template.xlsx"
    wb = openpyxl.Workbook()
    # --- daily sheet ---
    ws_days = wb.active
    assert ws_days is not None
    ws_days.title = "По_дням"
    ws_days["E5"] = "Период кампании"
    ws_days["F5"] = "01.07.2026 - 31.07.2026"
    ws_days["E6"] = "Период отчетности"
    ws_days["F6"] = "01.07.2026 - 05.07.2026"
    ws_days["E7"] = "Количество дней"
    ws_days["F7"] = 31
    ws_days["F10"] = "Период"
    ws_days["G10"] = "План"
    ws_days["H10"] = "Факт"
    for i, day in enumerate(range(6, 13)):
        row = 16 + i
        ws_days.cell(row=row, column=6).value = f"2026-07-{day:02d}"
        ws_days.cell(row=row, column=7).value = 100_000  # plan — must stay
        ws_days.cell(row=row, column=8).value = None

    # --- creatives sheet ---
    ws_cr = wb.create_sheet("Креативы")
    ws_cr["E5"] = "Период кампании"
    ws_cr["F5"] = "01.07.2026 - 31.07.2026"
    ws_cr["E10"] = "Креатив"
    ws_cr["G10"] = "План"
    ws_cr["H10"] = "Факт"
    for i, name in enumerate(("Creative_A", "Creative_B", "Creative_C")):
        row = 12 + i
        ws_cr.cell(row=row, column=5).value = name
        ws_cr.cell(row=row, column=7).value = 50_000
        ws_cr.cell(row=row, column=8).value = None

    wb.save(tmpl)
    wb.close()

    files = [
        tmpl,
        workspace / "cab_a_days.xlsx",
        workspace / "cab_b_days.xlsx",
        workspace / "cab_a_creatives.xlsx",
        workspace / "cab_b_creatives.xlsx",
    ]
    return {
        "template": str(tmpl),
        "files": files,
        "day_oracle": dict(DAY_ORACLE),
        "creative_oracle": dict(CREATIVE_ORACLE),
    }


def _write_day_source(path: Path, values: dict[str, int], *, headers: tuple[str, str]) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Report"
    ws["A1"] = headers[0]
    ws["B1"] = headers[1]
    for i, (day, val) in enumerate(sorted(values.items()), start=2):
        ws.cell(row=i, column=1).value = day
        ws.cell(row=i, column=2).value = val
    wb.save(path)
    wb.close()


def _write_creative_source(path: Path, values: dict[str, int], *, headers: tuple[str, str]) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Report"
    ws["A1"] = headers[0]
    ws["B1"] = headers[1]
    for i, (name, val) in enumerate(values.items(), start=2):
        ws.cell(row=i, column=1).value = name
        ws.cell(row=i, column=2).value = val
    wb.save(path)
    wb.close()
