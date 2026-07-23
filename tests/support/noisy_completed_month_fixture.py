"""Build a synthetic noisy workspace: completed Июнь + gap Июль + distractor June source.

Generates a fully deterministic template with two month-sheets and day-keyed
sources, so unit tests can verify the model fills only the target month (Июль)
and leaves the completed month (Июнь) untouched — without depending on any
real/commercial corpus.

Used by ``tests/test_noisy_completed_month.py``.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

__all__ = [
    "JUNE_SENTINEL_BASE",
    "JULY_SOURCE_BASE",
    "build_noisy_workspace",
    "june_h_fingerprint",
]

# Sentinel values written into the completed Июнь sheet. The test asserts these
# are unchanged after the Июль fill, proving the filler never touches a past month.
JUNE_SENTINEL_BASE = 910_000

# Synthetic day values for the two July sources. The test expects 2026-07-06 to
# sum to JULY_SOURCE_BASE["a"][6] + JULY_SOURCE_BASE["b"][6] = 100000 + 42997.
JULY_SOURCE_BASE: dict[str, dict[int, int]] = {
    "a": {6: 100_000, 7: 200_000},
    "b": {6: 42_997, 7: 30_000},
}


def build_noisy_workspace(workspace: Path) -> dict[str, Any]:
    """Prepare workspace with synthetic noisy template + July sources + June distractor.

    Returns metadata including ``june_fingerprint`` for untouched-oracle checks.
    """
    workspace.mkdir(parents=True, exist_ok=True)

    _write_day_source(workspace / "source_a.xlsx", JULY_SOURCE_BASE["a"])
    _write_day_source(workspace / "source_b.xlsx", JULY_SOURCE_BASE["b"])
    _write_june_distractor(workspace / "distractor_june.xlsx")

    tmpl = workspace / "template.xlsx"
    _write_synthetic_template(tmpl)

    fingerprint = june_h_fingerprint(tmpl, "Июнь")
    return {
        "template": str(tmpl),
        "june_fingerprint": fingerprint,
        "files": [
            tmpl,
            workspace / "source_a.xlsx",
            workspace / "source_b.xlsx",
            workspace / "distractor_june.xlsx",
        ],
    }


def june_h_fingerprint(path: Path | str, sheet: str = "Июнь") -> dict[str, int]:
    """Map H16..H45 → int values for untouched checks."""
    import openpyxl

    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[sheet]
        out: dict[str, int] = {}
        for row in range(16, 46):
            val = ws.cell(row=row, column=8).value
            out[f"H{row}"] = int(val) if isinstance(val, (int, float)) else -1
        return out
    finally:
        wb.close()


def _write_synthetic_template(path: Path) -> None:
    """Create a two-month template: completed Июнь + empty Июль.

    Both sheets mirror the media-report grid layout:
      F5/F6 = period strings, F12 = "Дата" header, rows 16+ = daily grid,
      column F = date, column H = fact (empty for Июль).
    """
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    # --- Июль (target month — empty fact column to be filled) ---
    july = wb.active
    assert july is not None
    july.title = "Июль"
    _layout_month_sheet(july, "01.07.2026", "31.07.2026", "01.07.2026", "05.07.2026")
    for day in range(1, 31):
        row = 15 + day
        july.cell(row=row, column=6).value = f"2026-07-{day:02d}"
        july.cell(row=row, column=7).value = 100_000  # plan — must stay
        july.cell(row=row, column=8).value = None  # fact — to be filled

    # --- Июнь (completed month — must be left untouched) ---
    june = wb.create_sheet("Июнь")
    _layout_month_sheet(june, "01.06.2026", "30.06.2026", "01.06.2026", "30.06.2026")
    for day in range(1, 31):
        row = 15 + day  # H16 = 01.06 … H45 = 30.06
        june.cell(row=row, column=6).value = f"2026-06-{day:02d}"
        june.cell(row=row, column=7).value = 100_000  # plan
        june.cell(row=row, column=8).value = JUNE_SENTINEL_BASE + day  # fact (filled)

    # Put Июнь first so brief lists the completed month before the target.
    wb.move_sheet(june, offset=-len(wb.sheetnames) + 1)
    # Ensure columns are wide enough for readability (non-functional).
    for ws in wb.worksheets:
        ws.column_dimensions[get_column_letter(6)].width = 16
        ws.column_dimensions[get_column_letter(8)].width = 12

    wb.save(path)
    wb.close()


def _layout_month_sheet(
    ws: Any,
    camp_start: str,
    camp_end: str,
    rep_start: str,
    rep_end: str,
) -> None:
    """Place the standard period / header labels on a month sheet."""
    ws["E5"] = "Период кампании"
    ws["F5"] = f"{camp_start} - {camp_end}"
    ws["E6"] = "Период отчетности"
    ws["F6"] = f"{rep_start} - {rep_end}"
    ws["F12"] = "Дата"
    ws["G12"] = "План"
    ws["H12"] = "Факт"


def _write_day_source(path: Path, day_values: dict[int, int]) -> None:
    """Small source workbook keyed by July day numbers (1..31)."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Report"
    ws["A1"] = "Дата"
    ws["B1"] = "Показы"
    for day, val in sorted(day_values.items()):
        ws.cell(row=day + 1, column=1).value = f"2026-07-{day:02d}"
        ws.cell(row=day + 1, column=2).value = val
    wb.save(path)
    wb.close()


def _write_june_distractor(path: Path) -> None:
    """Small source workbook with June-only impressions (must not feed July)."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Report"
    ws["A1"] = "Дата"
    ws["B1"] = "Показы"
    for day in range(1, 31):
        ws.cell(row=day + 1, column=1).value = f"2026-06-{day:02d}"
        ws.cell(row=day + 1, column=2).value = 50_000 + day
    wb.save(path)
    wb.close()
