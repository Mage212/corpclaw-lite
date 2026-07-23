"""Dual-type synthetic fill: daily sheet + creatives sheet from 4 sources.

Ported to clean FillPlan API (explicit plan, no suggest_fill_plan oracle).
"""

from __future__ import annotations

import shutil
import tempfile
from pathlib import Path

import openpyxl
import pytest
from support.dual_type_fixture import (  # type: ignore[import-not-found]
    CREATIVE_ORACLE,
    DAY_ORACLE,
    build_dual_type_workspace,
)

from corpclaw_lite.agent.fill_plan import (
    FillPeriodPlan,
    FillPlan,
    FillSheetPlan,
    FillSourceRef,
    apply_fill_plan,
)
from corpclaw_lite.agent.workbook_brief import (
    build_workbook_brief,
    format_files_brief_for_agent,
)


@pytest.fixture()
def dual_ws() -> Path:
    td = Path(tempfile.mkdtemp(prefix="dual_type_"))
    build_dual_type_workspace(td)
    yield td
    shutil.rmtree(td, ignore_errors=True)


def _dual_plan() -> FillPlan:
    return FillPlan(
        template="template.xlsx",
        output_path="target_report.xlsx",
        cutoff="2026-07-12",
        sheets=[
            FillSheetPlan(
                sheet="По_дням",
                date_column="F",
                value_column="H",
                aggregate="sum",
                match_mode="date",
                period=FillPeriodPlan(cell="F6", start_mode="month_start", format="dmy"),
                sources=[
                    FillSourceRef(
                        file="cab_a_days.xlsx",
                        sheet="Report",
                        date_column="A",
                        value_column="B",
                    ),
                    FillSourceRef(
                        file="cab_b_days.xlsx",
                        sheet="Report",
                        date_column="A",
                        value_column="B",
                    ),
                ],
            ),
            FillSheetPlan(
                sheet="Креативы",
                date_column="E",
                value_column="H",
                aggregate="sum",
                match_mode="key",
                sources=[
                    FillSourceRef(
                        file="cab_a_creatives.xlsx",
                        sheet="Report",
                        date_column="A",
                        value_column="B",
                    ),
                    FillSourceRef(
                        file="cab_b_creatives.xlsx",
                        sheet="Report",
                        date_column="A",
                        value_column="B",
                    ),
                ],
            ),
        ],
    )


def test_dual_brief_detects_day_and_key_sources(dual_ws: Path) -> None:
    files = list(dual_ws.glob("*.xlsx"))
    brief = build_workbook_brief(files)
    text = format_files_brief_for_agent(brief)
    assert "По_дням" in text
    assert "Креативы" in text
    assert "match_mode=key" in text
    assert "cab_a_days" in text
    assert "cab_a_creatives" in text


def test_dual_explicit_plan_wires_two_sheets() -> None:
    plan = _dual_plan()
    sheets = {s.sheet: s for s in plan.sheets}
    assert set(sheets) == {"По_дням", "Креативы"}
    assert sheets["По_дням"].match_mode == "date"
    assert sheets["Креативы"].match_mode == "key"
    day_files = {s.file for s in sheets["По_дням"].sources}
    cr_files = {s.file for s in sheets["Креативы"].sources}
    assert day_files == {"cab_a_days.xlsx", "cab_b_days.xlsx"}
    assert cr_files == {"cab_a_creatives.xlsx", "cab_b_creatives.xlsx"}


def test_dual_apply_fills_both_oracles(dual_ws: Path) -> None:
    results = apply_fill_plan(_dual_plan(), workspace=dual_ws)
    assert len(results) == 2
    assert not any("Error" in r for r in results)

    out = dual_ws / "target_report.xlsx"
    wb = openpyxl.load_workbook(out)
    try:
        days = wb["По_дням"]
        for i, day in enumerate(range(6, 13)):
            row = 16 + i
            key = f"2026-07-{day:02d}"
            assert int(days.cell(row=row, column=8).value or 0) == DAY_ORACLE[key]
            assert int(days.cell(row=row, column=7).value or 0) == 100_000
        assert "12.07.2026" in str(days["F6"].value or "")

        cr = wb["Креативы"]
        for i, name in enumerate(("Creative_A", "Creative_B", "Creative_C")):
            row = 12 + i
            assert cr.cell(row=row, column=5).value == name
            assert int(cr.cell(row=row, column=8).value or 0) == CREATIVE_ORACLE[name]
            assert int(cr.cell(row=row, column=7).value or 0) == 50_000
    finally:
        wb.close()
